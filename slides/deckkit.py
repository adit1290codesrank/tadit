"""Shared drawing helpers for the SIH deck slide builders (palette, fonts, shapes, icons)."""
import math
import re
from pathlib import Path

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ICONS = Path(__file__).resolve().parent / "icons"

NAVY, INK, INK2, MUTED = "0B2545", "1B2533", "3D4A5C", "6B7787"
BLUE, BOLT, SAFFRON, GREEN = "1565C0", "FFB400", "F37021", "138808"
RED, TEAL, VIOLET, AMBER = "D62828", "0F8B8D", "5B3FA8", "E08A00"
PANEL, LINE, WHITE = "F5F8FC", "D9E2EE", "FFFFFF"
LINK = "0B5CAD"

_LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
_TOKENS = re.compile(r"(\*\*|\[[^\]]+\]\([^)]+\))")

HEAD = "Segoe UI"          # ships with Windows; bold/black weights below
BLACK = "Segoe UI Black"
BODY = "Calibri"


def rgb(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------- helpers
class Slide:
    def __init__(self, slide):
        self.s = slide
        self.shapes = slide.shapes

    def box(self, x, y, w, h, fill=None, line=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE,
            radius=None, shadow=False, grad=None, grad_angle=0):
        sh = self.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        if grad:
            sh.fill.gradient()
            sh.fill.gradient_angle = grad_angle
            stops = sh.fill.gradient_stops
            stops[0].color.rgb, stops[0].position = rgb(grad[0]), 0.0
            stops[1].color.rgb, stops[1].position = rgb(grad[1]), 1.0
        elif fill:
            sh.fill.solid()
            sh.fill.fore_color.rgb = rgb(fill)
        else:
            sh.fill.background()
        if line:
            sh.line.color.rgb = rgb(line)
            sh.line.width = Pt(lw)
        else:
            sh.line.fill.background()
        if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = radius
        sh.shadow.inherit = False
        if shadow:
            add_shadow(sh)
        sh.text_frame.text = ""
        return sh

    def text(self, x, y, w, h, runs, size=9, color=INK2, font=BODY, bold=False, italic=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None, bold_color=None, space_after=0):
        """runs: str, or list of str (one paragraph each). Inside a str, **bold** and [label](url) links."""
        tb = self.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = None
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        paras = runs if isinstance(runs, list) else [runs]
        for i, ptxt in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            if spacing:
                p.line_spacing = spacing
            if space_after:
                p.space_after = Pt(space_after)
            in_bold = False
            for tok in _TOKENS.split(ptxt):
                if not tok:
                    continue
                if tok == "**":
                    in_bold = not in_bold
                    continue
                link = _LINK.fullmatch(tok)
                r = p.add_run()
                r.text = link.group(1) if link else tok
                f = r.font
                f.size = Pt(size)
                f.name = font
                f.italic = italic
                f.bold = bold or in_bold
                if link:
                    r.hyperlink.address = link.group(2)
                    f.color.rgb = rgb(LINK)
                    f.underline = True
                else:
                    f.color.rgb = rgb((bold_color or INK) if (in_bold and not bold) else color)
                _same_font_all_scripts(r, font)
        # PowerPoint's PDF export drops a hyperlink that is the very last run of a text box
        last = tf.paragraphs[-1].runs
        if last and last[-1].hyperlink.address:
            tail = tf.paragraphs[-1].add_run()
            tail.text = " "
            tail.font.size = Pt(size)
        return tb

    def image(self, path, x, y, w=None, h=None):
        return self.shapes.add_picture(str(path), Inches(x), Inches(y),
                                       Inches(w) if w else None, Inches(h) if h else None)

    def icon(self, name, x, y, size):
        return self.shapes.add_picture(str(ICONS / f"{name}.png"), Inches(x), Inches(y), Inches(size), Inches(size))

    def bubble(self, name, cx, cy, d, color, pad=0.2):
        self.box(cx - d / 2, cy - d / 2, d, d, fill=color, shape=MSO_SHAPE.OVAL)
        s = d * (1 - 2 * pad)
        self.icon(name, cx - s / 2, cy - s / 2, s)

    def line(self, pts, color, width=1.5):
        fb = self.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]), scale=1.0)
        fb.add_line_segments([(Inches(a), Inches(b)) for a, b in pts[1:]], close=False)
        sh = fb.convert_to_shape()
        sh.fill.background()
        sh.line.color.rgb = rgb(color)
        sh.line.width = Pt(width)
        sh.shadow.inherit = False
        return sh

    def dot(self, cx, cy, color, d=0.085):
        self.box(cx - d / 2, cy - d / 2, d, d, fill=WHITE, line=color, lw=1.5, shape=MSO_SHAPE.OVAL)


def add_shadow(sh, blur=4, dist=1.5, alpha=14):
    spPr = sh._element.spPr
    for old in spPr.findall(qn("a:effectLst")):
        spPr.remove(old)
    eff = etree.SubElement(spPr, qn("a:effectLst"))
    sd = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(Pt(blur)), dist=str(Pt(dist)),
                          dir="5400000", algn="t", rotWithShape="0")
    clr = etree.SubElement(sd, qn("a:srgbClr"), val=NAVY)
    etree.SubElement(clr, qn("a:alpha"), val=str(alpha * 1000))


def _same_font_all_scripts(run, font):
    latin = run._r.get_or_add_rPr().find(qn("a:latin"))
    if latin is None:
        return
    prev = latin
    for tag in ("a:ea", "a:cs"):
        el = etree.Element(qn(tag), typeface=font)
        prev.addnext(el)
        prev = el


def reset_slide(slide):
    """Remove the template's grey instruction text box (builders start from the pristine template)."""
    for sh in list(slide.shapes):
        if sh.name == "TextBox 8" and sh.has_text_frame and sh.top > Inches(2):
            sh._element.getparent().remove(sh._element)


def find_slide(prs, title):
    return next(s for s in prs.slides
                if any(sh.has_text_frame and title in sh.text_frame.text for sh in s.shapes))


def set_title(slide, text, size=None):
    """Replace the title placeholder's text, keeping its template formatting."""
    tf = slide.shapes.title.text_frame
    runs = tf.paragraphs[0].runs
    runs[0].text = text
    slide.shapes.title.top = Inches(-0.05)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    if size:
        runs[0].font.size = Pt(size)
