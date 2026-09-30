"""Build the whole SIH idea deck from the untouched template.

Usage (from the repo root):
    python slides/build_deck.py                       # updates SIH-PPT.pptx
    python slides/build_deck.py --pdf                 # ... and SIH-PPT.pdf, via PowerPoint
    python slides/build_deck.py --team-name X --team-id Y

Every run starts again from slides/template/, so it is safe to re-run after changing any number.
Manual edits made in PowerPoint are overwritten: put them in the builders instead.
"""
import argparse
import importlib.util
import subprocess
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))

from pptx import Presentation
from pptx.util import Pt

from deckkit import BLUE, NAVY, find_slide, rgb  # noqa: E402

TEMPLATE = HERE / "template" / "SIH2026-IDEA-Presentation-Format.pptx"
OUT = ROOT / "SIH-PPT.pptx"


def load(rel):
    spec = importlib.util.spec_from_file_location(Path(rel).stem, HERE / rel)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def title_page(slide, team_id, team_name):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        tf = sh.text_frame
        if "TITLE PAGE" in tf.text:
            p = [p for p in tf.paragraphs if "TITLE PAGE" in p.text][0]
            r0 = p.runs[0]
            r0.text = "TADIT "
            r0.font.color.rgb = rgb(NAVY)
            r1 = p.add_run()
            r1.text = "(तड़ित)"
            r1.font.name = "Nirmala UI"
            r1.font.size = Pt(28)
            r1.font.bold = True
            r1.font.color.rgb = rgb("F37021")
            for extra in p.runs[1:-1]:
                extra._r.getparent().remove(extra._r)
            q = tf.add_paragraph()
            q.alignment = p.alignment
            rq = q.add_run()
            rq.text = "AI nowcasting of thunderstorms & lightning, 0–6 hours ahead"
            rq.font.size = Pt(16)
            rq.font.italic = True
            rq.font.name = "Times New Roman"
            rq.font.color.rgb = rgb("3D4A5C")
        if "Problem Statement ID" in tf.text:
            for p in tf.paragraphs:
                runs = p.runs
                if not runs:
                    continue
                label = runs[0].text
                if len(runs) > 1:
                    for r in runs[1:]:
                        r.font.color.rgb = rgb(BLUE)
                if label.startswith("Team ID") and team_id:
                    r = p.add_run()
                    r.text = f" {team_id}"
                    r.font.bold = False
                    r.font.color.rgb = rgb(BLUE)
                if label.startswith("Team Name") and team_name:
                    r = p.add_run()
                    r.text = f" – {team_name}"
                    r.font.bold = False
                    r.font.color.rgb = rgb(BLUE)


def team_ovals(prs, team_name):
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip() == "Your Team Name":
                runs = sh.text_frame.paragraphs[0].runs
                runs[0].text = team_name
                for r in runs[1:]:
                    r._r.getparent().remove(r._r)
                runs[0].font.size = Pt(14 if len(team_name) <= 12 else 11)
                runs[0].font.bold = True


def drop_instructions(prs):
    sld_ids = prs.slides._sldIdLst
    for sld_id, slide in zip(list(sld_ids), list(prs.slides)):
        if any(sh.has_text_frame and "IMPORTANT INSTRUCTIONS" in sh.text_frame.text for sh in slide.shapes):
            prs.part.drop_rel(sld_id.rId)
            sld_ids.remove(sld_id)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(OUT))
    ap.add_argument("--team-name", default="")
    ap.add_argument("--team-id", default="")
    ap.add_argument("--pdf", action="store_true")
    a = ap.parse_args()

    prs = Presentation(str(TEMPLATE))
    title_page(prs.slides[0], a.team_id, a.team_name)
    load("solution/build_slide2.py").build(find_slide(prs, "IDEA TITLE"))
    load("technical/build_slide3.py").build(find_slide(prs, "TECHNICAL APPROACH"))
    load("feasibility_viability/build_slide4.py").build(find_slide(prs, "FEASIBILITY AND VIABILITY"))
    load("impact_benefits/build_slide5.py").build(find_slide(prs, "IMPACT AND BENEFITS"))
    load("references/build_slide6.py").build(find_slide(prs, "RESEARCH"))
    drop_instructions(prs)
    if a.team_name:
        team_ovals(prs, a.team_name)
    prs.save(a.out)
    print("wrote", a.out)

    if a.pdf:
        pdf = str(Path(a.out).with_suffix(".pdf"))
        subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(HERE / "render.ps1"),
                        "-Deck", str(Path(a.out).resolve()), "-Slide", "1",
                        "-Png", str(HERE / "_cover.png"), "-Pdf", pdf], check=True)
        (HERE / "_cover.png").unlink(missing_ok=True)
        print("wrote", pdf)


if __name__ == "__main__":
    main()
